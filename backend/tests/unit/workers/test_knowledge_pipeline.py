"""离线文档 pipeline 的输出与共享向量版本契约。"""

from types import SimpleNamespace
from unittest.mock import AsyncMock

from app.modules.knowledge.version import embedding_version
from app.workers.knowledge.config import WorkerSettings


async def test_pipeline_tags_every_chunk_with_vector_space_version(monkeypatch):
    from app.workers.knowledge.pipeline import KnowledgePipeline

    settings = WorkerSettings(_env_file=None)
    pipeline = KnowledgePipeline(settings)
    block = SimpleNamespace(
        ordinal=0, heading="标题", content="正文", page_start=None, page_end=None, kind="text"
    )
    monkeypatch.setattr(pipeline, "_prepare", lambda suffix, data: [block])
    pipeline._embed = AsyncMock(return_value=[[1.0, 0.0]])
    chunks = await pipeline.process(".txt", b"text")
    assert chunks[0]["embedding_version"] == embedding_version(settings)
    pipeline._embed.assert_awaited_once_with(["标题\n正文"])


async def test_text_document_runs_real_loading_cleaning_and_splitting():
    from app.workers.knowledge.pipeline import KnowledgePipeline

    settings = WorkerSettings(_env_file=None)
    pipeline = KnowledgePipeline(settings)

    async def embed(texts):
        return [[1.0, 0.0] for _ in texts]

    pipeline._embed = AsyncMock(side_effect=embed)
    chunks = await pipeline.process(".txt", "质押规则：奖励按周期发放。\n\n用户可随时查看状态。".encode())
    assert chunks
    assert "奖励按周期发放" in "\n".join(chunk["content"] for chunk in chunks)
    assert {chunk["embedding_version"] for chunk in chunks} == {embedding_version(settings)}


async def test_docx_parsing_preserves_heading_and_content():
    from io import BytesIO

    from docx import Document

    from app.workers.knowledge.pipeline import KnowledgePipeline

    document = Document()
    document.add_heading("质押指南", level=1)
    document.add_paragraph("奖励按周期发放，用户可以查看处理状态。")
    stream = BytesIO()
    document.save(stream)
    pipeline = KnowledgePipeline(WorkerSettings(_env_file=None))
    pipeline._embed = AsyncMock(side_effect=lambda texts: [[1.0, 0.0] for _ in texts])
    chunks = await pipeline.process(".docx", stream.getvalue())
    assert "奖励按周期发放" in "\n".join(chunk["content"] for chunk in chunks)
    assert any("质押指南" in (chunk["heading"] or "") for chunk in chunks)
